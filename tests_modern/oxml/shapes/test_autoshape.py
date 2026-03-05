from __future__ import annotations

import pytest
from syrupy.assertion import SnapshotAssertion

from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml import parse_xml
from pptx.oxml.shapes.autoshape import CT_Shape
from pptx.oxml.shapes.shared import ST_Direction, ST_PlaceholderSize


def test_prst_geom_gd_lst() -> None:
    prst_geom = parse_xml(
        '<a:prstGeom xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'prst="chevron"><a:avLst><a:gd name="adj1" fmla="val 111"/>'
        '<a:gd name="adj2" fmla="val 222"/></a:avLst></a:prstGeom>'
    )

    gd_vals = [(gd.name, gd.fmla) for gd in prst_geom.gd_lst]

    assert gd_vals == [("adj1", "val 111"), ("adj2", "val 222")]


def test_prst_geom_rewrite_guides(snapshot: SnapshotAssertion) -> None:
    prst_geom = parse_xml(
        '<a:prstGeom xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'prst="chevron"><a:avLst><a:gd name="adj6" fmla="val 666"/></a:avLst></a:prstGeom>'
    )

    prst_geom.rewrite_guides([("adj1", 111), ("adj2", 222)])

    assert str(prst_geom.xml) == snapshot


def test_shape_new_autoshape_sp(snapshot: SnapshotAssertion) -> None:
    sp = CT_Shape.new_autoshape_sp(9, "Rounded Rectangle 8", "roundRect", 111, 222, 333, 444)

    assert str(sp.xml) == snapshot


@pytest.mark.parametrize(
    ("id_", "name", "ph_type", "orient", "sz", "idx"),
    [
        (
            2,
            "Title 1",
            PP_PLACEHOLDER.CENTER_TITLE,
            ST_Direction.HORZ,
            ST_PlaceholderSize.FULL,
            0,
        ),
        (
            4,
            "Vertical Subtitle 3",
            PP_PLACEHOLDER.SUBTITLE,
            ST_Direction.VERT,
            ST_PlaceholderSize.FULL,
            1,
        ),
        (
            7,
            "Footer Placeholder 6",
            PP_PLACEHOLDER.FOOTER,
            ST_Direction.HORZ,
            ST_PlaceholderSize.QUARTER,
            11,
        ),
    ],
)
def test_shape_new_placeholder_sp(
    id_: int,
    name: str,
    ph_type: PP_PLACEHOLDER,
    orient: ST_Direction | str,
    sz: ST_PlaceholderSize | str,
    idx: int,
    snapshot: SnapshotAssertion,
) -> None:
    sp = CT_Shape.new_placeholder_sp(id_, name, ph_type, orient, sz, idx)

    assert str(sp.xml) == snapshot


def test_shape_new_textbox_sp(snapshot: SnapshotAssertion) -> None:
    sp = CT_Shape.new_textbox_sp(9, "TextBox 8", 111, 222, 333, 444)

    assert str(sp.xml) == snapshot


@pytest.mark.parametrize(
    ("xml", "expected"),
    [
        (
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<p:nvSpPr><p:cNvSpPr/></p:nvSpPr><p:spPr><a:prstGeom/></p:spPr></p:sp>',
            True,
        ),
        (
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<p:nvSpPr><p:nvPr><p:ph/></p:nvPr></p:nvSpPr><p:spPr/></p:sp>',
            False,
        ),
        (
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<p:nvSpPr><p:cNvSpPr txBox="1"/></p:nvSpPr><p:spPr><a:prstGeom/></p:spPr></p:sp>',
            False,
        ),
    ],
)
def test_shape_is_autoshape(xml: str, expected: bool) -> None:
    sp = parse_xml(xml)

    assert sp.is_autoshape is expected


@pytest.mark.parametrize(
    ("xml", "expected"),
    [
        (
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<p:nvSpPr><p:cNvSpPr/></p:nvSpPr><p:spPr><a:prstGeom/></p:spPr></p:sp>',
            False,
        ),
        (
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
            '<p:nvSpPr><p:cNvSpPr/><p:nvPr><p:ph/></p:nvPr></p:nvSpPr><p:spPr/></p:sp>',
            False,
        ),
        (
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<p:nvSpPr><p:cNvSpPr txBox="1"/></p:nvSpPr><p:spPr><a:prstGeom/></p:spPr></p:sp>',
            True,
        ),
    ],
)
def test_shape_is_textbox(xml: str, expected: bool) -> None:
    sp = parse_xml(xml)

    assert sp.is_textbox is expected


def test_shape_add_path(snapshot: SnapshotAssertion) -> None:
    sp = parse_xml(
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:spPr><a:custGeom/></p:spPr></p:sp>'
    )

    path = sp.add_path(100, 200)

    assert path.tag.endswith("path")
    assert str(sp.xml) == snapshot


def test_shape_add_path_raises_if_not_freeform() -> None:
    sp = parse_xml(
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:spPr/></p:sp>'
    )

    with pytest.raises(ValueError, match="shape must be freeform"):
        sp.add_path(100, 200)


def test_shape_get_or_add_ln(snapshot: SnapshotAssertion) -> None:
    sp = parse_xml(
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:spPr/></p:sp>'
    )

    assert sp.ln is None

    ln = sp.get_or_add_ln()

    assert sp.ln is ln
    assert str(sp.xml) == snapshot


def test_shape_prst_value() -> None:
    sp = parse_xml(
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:spPr><a:prstGeom prst="rect"/></p:spPr></p:sp>'
    )

    assert sp.prst == MSO_AUTO_SHAPE_TYPE.RECTANGLE


def test_shape_prst_none() -> None:
    sp = parse_xml(
        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:spPr/></p:sp>'
    )

    assert sp.prst is None


def test_path2d_operations(snapshot: SnapshotAssertion) -> None:
    path = parse_xml('<a:path xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>')

    path.add_moveTo(10, 20)
    path.add_lnTo(30, 40)
    path.add_close()

    assert str(path.xml) == snapshot
