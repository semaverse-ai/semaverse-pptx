from __future__ import annotations

import pytest
from syrupy.assertion import SnapshotAssertion

from pptx.dml.fill import FillFormat
from pptx.enum.text import MSO_ANCHOR
from pptx.table import Table, _Cell
from pptx.util import Inches
from tests.stubs import GraphicFrameProxy
from tests.text.factories import table
from tests.xml_utils import serialize_xml


def _table(xml_body: bytes) -> object:
    return table(xml_body)


def test_table_cell_access() -> None:
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='100'/><a:gridCol w='200'/></a:tblGrid>"
        b"<a:tr h='300'><a:tc id='00'/><a:tc id='01'/></a:tr>"
        b"<a:tr h='400'><a:tc id='10'/><a:tc id='11'/></a:tr>"
    )
    table_obj = Table(tbl, GraphicFrameProxy())

    assert table_obj.cell(0, 0)._tc.get("id") == "00"
    assert table_obj.cell(0, 1)._tc.get("id") == "01"
    assert table_obj.cell(1, 0)._tc.get("id") == "10"
    assert table_obj.cell(1, 1)._tc.get("id") == "11"


def test_table_iter_cells() -> None:
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='100'/><a:gridCol w='200'/></a:tblGrid>"
        b"<a:tr h='300'><a:tc id='00'/><a:tc id='01'/></a:tr>"
        b"<a:tr h='400'><a:tc id='10'/><a:tc id='11'/></a:tr>"
    )
    table_obj = Table(tbl, GraphicFrameProxy())

    cell_ids = [cell._tc.get("id") for cell in table_obj.iter_cells()]

    assert cell_ids == ["00", "01", "10", "11"]


def test_table_columns_rows_and_resize_notifications() -> None:
    graphic_frame = GraphicFrameProxy()
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='100'/><a:gridCol w='200'/></a:tblGrid>"
        b"<a:tr h='300'><a:tc/><a:tc/></a:tr>"
        b"<a:tr h='400'><a:tc/><a:tc/></a:tr>"
    )
    table_obj = Table(tbl, graphic_frame)

    assert len(table_obj.columns) == 2
    assert len(table_obj.rows) == 2
    assert table_obj.columns[0].width == 100
    assert table_obj.rows[0].height == 300

    table_obj.columns[0].width = 150
    table_obj.rows[0].height = 350

    assert graphic_frame.width == 350
    assert graphic_frame.height == 750


def test_table_part_property() -> None:
    part = object()
    table_obj = Table(
        _table(b"<a:tblGrid/><a:tr h='100'><a:tc/></a:tr>"),
        GraphicFrameProxy(part=part),
    )

    assert table_obj.part is part


@pytest.mark.parametrize(
    ("prop_name", "xml_attr"),
    [
        ("first_row", b"firstRow"),
        ("first_col", b"firstCol"),
        ("last_row", b"lastRow"),
        ("last_col", b"lastCol"),
        ("horz_banding", b"bandRow"),
        ("vert_banding", b"bandCol"),
    ],
)
def test_table_boolean_properties(prop_name: str, xml_attr: bytes) -> None:
    tbl = _table(b"<a:tblPr " + xml_attr + b"='1'/>")
    table_obj = Table(tbl, GraphicFrameProxy())

    assert getattr(table_obj, prop_name) is True

    setattr(table_obj, prop_name, False)

    assert getattr(table_obj, prop_name) is False


def test_cell_equality() -> None:
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='100'/></a:tblGrid>"
        b"<a:tr h='100'><a:tc/><a:tc/></a:tr>"
    )
    table_obj = Table(tbl, GraphicFrameProxy())
    cell = table_obj.cell(0, 0)
    cell_same_tc = _Cell(cell._tc, table_obj)
    cell_other = table_obj.cell(0, 1)

    assert cell == cell_same_tc
    assert cell != cell_other
    assert cell != object()


def test_cell_fill() -> None:
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='100'/></a:tblGrid>"
        b"<a:tr h='100'><a:tc><a:tcPr/></a:tc></a:tr>"
    )
    table_obj = Table(tbl, GraphicFrameProxy())

    assert isinstance(table_obj.cell(0, 0).fill, FillFormat)


def test_cell_margin_and_vertical_anchor() -> None:
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='100'/></a:tblGrid>"
        b"<a:tr h='100'><a:tc>"
        b"<a:tcPr marL='10' marR='20' marT='30' marB='40' anchor='ctr'/>"
        b"</a:tc></a:tr>"
    )
    cell = Table(tbl, GraphicFrameProxy()).cell(0, 0)

    assert cell.margin_left == 10
    assert cell.margin_right == 20
    assert cell.margin_top == 30
    assert cell.margin_bottom == 40
    assert cell.vertical_anchor == MSO_ANCHOR.MIDDLE

    cell.margin_left = Inches(1)
    cell.margin_right = Inches(2)
    cell.margin_top = Inches(3)
    cell.margin_bottom = Inches(4)
    cell.vertical_anchor = MSO_ANCHOR.TOP

    assert cell.margin_left == 914400
    assert cell.margin_right == 1828800
    assert cell.margin_top == 2743200
    assert cell.margin_bottom == 3657600
    assert cell.vertical_anchor == MSO_ANCHOR.TOP


def test_cell_margin_raises_on_invalid_type() -> None:
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='100'/></a:tblGrid>"
        b"<a:tr h='100'><a:tc/></a:tr>"
    )
    cell = Table(tbl, GraphicFrameProxy()).cell(0, 0)

    with pytest.raises(TypeError):
        setattr(cell, "margin_left", "12")


def test_cell_text_property_round_trip() -> None:
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='100'/></a:tblGrid>"
        b"<a:tr h='100'><a:tc><a:txBody><a:bodyPr/><a:lstStyle/>"
        b"<a:p><a:r><a:t>Hello</a:t></a:r></a:p>"
        b"</a:txBody></a:tc></a:tr>"
    )
    cell = Table(tbl, GraphicFrameProxy()).cell(0, 0)

    assert cell.text == "Hello"

    cell.text = "World"

    assert cell.text == "World"


def test_cell_merge_and_split(snapshot: SnapshotAssertion) -> None:
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='100'/><a:gridCol w='100'/></a:tblGrid>"
        b"<a:tr h='100'>"
        b"<a:tc><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>A</a:t></a:r></a:p></a:txBody></a:tc>"
        b"<a:tc><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>B</a:t></a:r></a:p></a:txBody></a:tc>"
        b"</a:tr>"
        b"<a:tr h='100'>"
        b"<a:tc><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>C</a:t></a:r></a:p></a:txBody></a:tc>"
        b"<a:tc><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>D</a:t></a:r></a:p></a:txBody></a:tc>"
        b"</a:tr>"
    )
    table_obj = Table(tbl, GraphicFrameProxy())
    origin = table_obj.cell(0, 0)
    other = table_obj.cell(1, 1)

    origin.merge(other)

    assert origin.is_merge_origin is True
    assert origin.span_height == 2
    assert origin.span_width == 2
    assert table_obj.cell(0, 1).is_spanned is True
    assert table_obj.cell(1, 0).is_spanned is True
    assert table_obj.cell(1, 1).is_spanned is True
    assert serialize_xml(tbl) == snapshot(name="merged")

    origin.split()

    assert origin.is_merge_origin is False
    assert origin.span_height == 1
    assert origin.span_width == 1
    assert table_obj.cell(1, 1).is_spanned is False
    assert serialize_xml(tbl) == snapshot(name="split")


def test_cell_merge_raises_on_other_table() -> None:
    table_a = Table(
        _table(b"<a:tblGrid><a:gridCol w='1'/></a:tblGrid><a:tr h='1'><a:tc/></a:tr>"),
        GraphicFrameProxy(),
    )
    table_b = Table(
        _table(b"<a:tblGrid><a:gridCol w='1'/></a:tblGrid><a:tr h='1'><a:tc/></a:tr>"),
        GraphicFrameProxy(),
    )

    with pytest.raises(ValueError, match="different table"):
        table_a.cell(0, 0).merge(table_b.cell(0, 0))


def test_cell_merge_raises_on_existing_merge() -> None:
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='1'/><a:gridCol w='1'/></a:tblGrid>"
        b"<a:tr h='1'><a:tc/><a:tc/></a:tr>"
        b"<a:tr h='1'><a:tc/><a:tc/></a:tr>"
    )
    table_obj = Table(tbl, GraphicFrameProxy())

    table_obj.cell(0, 0).merge(table_obj.cell(1, 1))

    with pytest.raises(ValueError, match="contains one or more merged"):
        table_obj.cell(0, 0).merge(table_obj.cell(0, 1))


def test_cell_split_raises_on_non_merge_origin() -> None:
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='1'/></a:tblGrid>"
        b"<a:tr h='1'><a:tc/></a:tr>"
    )
    table_obj = Table(tbl, GraphicFrameProxy())

    with pytest.raises(ValueError, match="not a merge-origin"):
        table_obj.cell(0, 0).split()


def test_cell_collection_index_and_iter() -> None:
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='1'/><a:gridCol w='1'/></a:tblGrid>"
        b"<a:tr h='1'><a:tc id='0'/><a:tc id='1'/></a:tr>"
    )
    row = Table(tbl, GraphicFrameProxy()).rows[0]

    assert len(row.cells) == 2
    assert [cell._tc.get("id") for cell in row.cells] == ["0", "1"]
    assert row.cells[0]._tc.get("id") == "0"

    with pytest.raises(IndexError, match="cell index"):
        _ = row.cells[2]


def test_column_collection_index_error() -> None:
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='1'/></a:tblGrid>"
        b"<a:tr h='1'><a:tc/></a:tr>"
    )
    cols = Table(tbl, GraphicFrameProxy()).columns

    with pytest.raises(IndexError, match="column index"):
        _ = cols[1]


def test_row_collection_index_error() -> None:
    tbl = _table(
        b"<a:tblGrid><a:gridCol w='1'/></a:tblGrid>"
        b"<a:tr h='1'><a:tc/></a:tr>"
    )
    rows = Table(tbl, GraphicFrameProxy()).rows

    with pytest.raises(IndexError, match="row index"):
        _ = rows[1]
