from __future__ import annotations

from typing import Any

import pytest

from pptx.enum.shapes import MSO_CONNECTOR_TYPE, PP_PLACEHOLDER
from pptx.oxml import parse_xml


@pytest.fixture
def sp_tree() -> Any:
    return parse_xml(
        '<p:spTree xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'
    )


def test_group_shape_add_autoshape(sp_tree: Any) -> None:
    sp_tree.add_autoshape(42, "name", "rect", 9, 8, 7, 6)

    assert sp_tree[-1].tag.endswith("sp")
    assert sp_tree[-1].nvSpPr.cNvPr.id == 42


def test_group_shape_add_cxn_sp(sp_tree: Any) -> None:
    sp_tree.add_cxnSp(42, "name", MSO_CONNECTOR_TYPE.STRAIGHT, 9, 8, 7, 6, False, True)

    assert sp_tree[-1].tag.endswith("cxnSp")
    assert sp_tree[-1].nvCxnSpPr.cNvPr.id == 42


def test_group_shape_add_freeform_sp(sp_tree: Any) -> None:
    sp_tree.add_freeform_sp(9, 8, 7, 6)

    assert sp_tree[-1].tag.endswith("sp")
    assert sp_tree[-1].spPr.custGeom is not None


def test_group_shape_add_grp_sp(sp_tree: Any) -> None:
    grp_sp = sp_tree.add_grpSp()

    assert grp_sp.tag.endswith("grpSp")


def test_group_shape_add_pic(sp_tree: Any) -> None:
    sp_tree.add_pic(42, "name", "desc", "rId6", 6, 7, 8, 9)

    assert sp_tree[-1].tag.endswith("pic")


def test_group_shape_add_placeholder(sp_tree: Any) -> None:
    sp_tree.add_placeholder(42, "name", PP_PLACEHOLDER.OBJECT, "horz", "full", 24)

    assert sp_tree[-1].tag.endswith("sp")


def test_group_shape_add_table(sp_tree: Any) -> None:
    sp_tree.add_table(42, "name", 2, 3, 5, 4, 3, 2)

    assert sp_tree[-1].tag.endswith("graphicFrame")


def test_group_shape_add_textbox(sp_tree: Any) -> None:
    sp_tree.add_textbox(42, "name", 3, 4, 5, 6)

    assert sp_tree[-1].tag.endswith("sp")


def test_group_shape_recalculate_extents() -> None:
    grp_sp = parse_xml(
        '<p:grpSp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
        '<a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
        '<p:sp><p:spPr><a:xfrm><a:off x="100" y="200"/><a:ext cx="300" cy="400"/>'
        "</a:xfrm></p:spPr></p:sp>"
        '<p:sp><p:spPr><a:xfrm><a:off x="150" y="250"/><a:ext cx="300" cy="400"/>'
        "</a:xfrm></p:spPr></p:sp>"
        "</p:grpSp>"
    )

    parent = parse_xml(
        '<p:spTree xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'
    )
    parent.append(grp_sp)

    grp_sp.recalculate_extents()

    assert grp_sp.grpSpPr.xfrm.chOff.x == 100
    assert grp_sp.grpSpPr.xfrm.chOff.y == 200
    assert grp_sp.grpSpPr.xfrm.chExt.cx == 350
    assert grp_sp.grpSpPr.xfrm.chExt.cy == 450


def test_group_shape_child_extents() -> None:
    grp_sp = parse_xml(
        '<p:grpSp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:sp><p:spPr><a:xfrm><a:off x="100" y="200"/><a:ext cx="300" cy="400"/>'
        "</a:xfrm></p:spPr></p:sp>"
        '<p:sp><p:spPr><a:xfrm><a:off x="150" y="250"/><a:ext cx="300" cy="400"/>'
        "</a:xfrm></p:spPr></p:sp>"
        "</p:grpSp>"
    )

    x, y, cx, cy = grp_sp._child_extents

    assert (x, y, cx, cy) == (100, 200, 350, 450)


def test_group_shape_child_extents_empty() -> None:
    grp_sp = parse_xml(
        '<p:grpSp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'
    )

    x, y, cx, cy = grp_sp._child_extents

    assert (x, y, cx, cy) == (0, 0, 0, 0)


def test_group_shape_max_shape_id() -> None:
    sp_tree = parse_xml(
        '<p:spTree xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<p:sp><p:nvSpPr><p:cNvPr id="5"/></p:nvSpPr></p:sp>'
        '<p:sp><p:nvSpPr><p:cNvPr id="2"/></p:nvSpPr></p:sp>'
        "</p:spTree>"
    )

    assert sp_tree.max_shape_id == 5


def test_group_shape_max_shape_id_empty() -> None:
    sp_tree = parse_xml(
        '<p:spTree xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>'
    )

    assert sp_tree.max_shape_id == 0
